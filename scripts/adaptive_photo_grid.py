# -*- coding: utf-8 -*-
"""
python-docx-photo-grid 技能 — 可重用代码模块
================================================
从 python-docx-photo-grid 技能抽取的核心函数，可直接 import 使用。

用法：
    import sys
    sys.path.insert(0, r"<SKILLS_DIR>/python-docx-photo-grid/scripts")
    from adaptive_photo_grid import (
        setup_a4_page, GRID_COLS, MAX_ROWS,
        compress_image, add_photo_grid, verify_docx
    )

依赖：python-docx, Pillow（managed venv：envs/default）
"""
import os
import glob
import hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from PIL import Image

# ===== A4 纵向几何 =====
PAGE_W, PAGE_H = 21.0, 29.7
MARGIN = 2.54
CONTENT_W = PAGE_W - 2 * MARGIN      # 15.92
CONTENT_H = PAGE_H - 2 * MARGIN      # 24.62

# ===== 网格参数 =====
GRID_COLS = 2
MAX_ROWS = 3
COL_W = CONTENT_W / GRID_COLS        # 7.96
IMG_MAX_W = COL_W - 0.96             # 7.00
RESERVE_H = 2.2
PAGE_BUDGET = CONTENT_H - RESERVE_H  # 22.42
CAPTION_H = 0.9
CELL_PAD = 0.5
SAFETY = 0.15
ABS_MAX_IMG_H = 9.0
MIN_IMG_H = 3.0

FONT_CN = "PMingLiU"
FONT_EN = "Times New Roman"


def set_run_font(run, size=12, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = FONT_EN
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), FONT_CN)


def set_table_borders(table, color="000000", sz="6"):
    tbl = table._element
    tblPr = tbl.tblPr
    borders = tblPr.makeelement(qn('w:tblBorders'), {})
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = borders.makeelement(qn(f'w:{edge}'), {
            qn('w:val'): 'single', qn('w:sz'): sz,
            qn('w:space'): '0', qn('w:color'): color
        })
        borders.append(b)
    tblPr.append(borders)


def fix_table_layout(table, col_w):
    table.autofit = False
    tblPr = table._element.tblPr
    layout = tblPr.makeelement(qn('w:tblLayout'), {qn('w:type'): 'fixed'})
    tblPr.append(layout)
    for row in table.rows:
        for cell in row.cells:
            cell.width = Cm(col_w)


def compress_image(src, dst, max_width_cm=12, dpi=200):
    img = Image.open(src)
    max_w = int(max_width_cm * dpi / 2.54)
    if img.width > max_w:
        ratio = max_w / img.width
        img = img.resize((max_w, int(img.height * ratio)), Image.LANCZOS)
    img.save(dst, "JPEG", quality=85, dpi=(dpi, dpi))
    return dst


def rendered_size(path, max_w, max_h):
    img = Image.open(path)
    ar = img.width / img.height
    w = max_w; h = w / ar
    if h > max_h:
        h = max_h; w = h * ar
    return w, h


def add_picture_fit(paragraph, path, max_w_cm, max_h_cm):
    w, h = rendered_size(path, max_w_cm, max_h_cm)
    run = paragraph.add_run()
    run.add_picture(path, width=Cm(w), height=Cm(h))
    return run


def setup_a4_page(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(PAGE_W)
    sec.page_height = Cm(PAGE_H)
    sec.top_margin = sec.bottom_margin = Cm(MARGIN)
    sec.left_margin = sec.right_margin = Cm(MARGIN)
    style = doc.styles['Normal']
    style.font.name = FONT_EN
    style.font.size = Pt(12)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), FONT_CN)


def add_photo_grid(doc, items, caption_size=9):
    """自适应 2 栏相片网格（核心算法）。
    items: [(image_path, caption_text), ...]"""
    per_page = MAX_ROWS * GRID_COLS
    pages = [items[i:i + per_page] for i in range(0, len(items), per_page)]
    counts = []
    for pi, page in enumerate(pages):
        rows = (len(page) + GRID_COLS - 1) // GRID_COLS
        h_allow = (PAGE_BUDGET / rows) - CAPTION_H - CELL_PAD - SAFETY
        h_allow = max(min(h_allow, ABS_MAX_IMG_H), MIN_IMG_H)

        table = doc.add_table(rows=rows, cols=GRID_COLS)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        fix_table_layout(table, COL_W)

        for idx, (img_path, caption) in enumerate(page):
            ri, ci = divmod(idx, GRID_COLS)
            cell = table.cell(ri, ci)
            p_cap = cell.paragraphs[0]
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(p_cap.add_run(caption), size=caption_size,
                         bold=True, color="595959")
            p_img = cell.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if img_path and os.path.exists(img_path):
                add_picture_fit(p_img, img_path, IMG_MAX_W, h_allow)
            else:
                set_run_font(p_img.add_run("[相片缺失]"), size=9, color="C00000")

        for idx in range(len(page), rows * GRID_COLS):
            ri, ci = divmod(idx, GRID_COLS)
            table.cell(ri, ci).text = ""
        set_table_borders(table)
        counts.append(len(page))
        if pi < len(pages) - 1:
            doc.add_page_break()
    return counts


def md5_dedup_photos(file_list):
    """对相片列表做 MD5 去重，保持原顺序。"""
    seen = set()
    result = []
    for p in file_list:
        with open(p, 'rb') as fp:
            m = hashlib.md5(fp.read()).hexdigest()
        if m in seen:
            continue
        seen.add(m)
        result.append(p)
    return result


def verify_docx(path, expect_imgs=None, verbose=True):
    """验证自适应网格 docx 的结构完整性。"""
    EMU = 360000
    d = Document(path)
    tables = d.tables
    imgs = [r for r in d.part.rels.values() if "image" in r.reltype]

    photo_pages = []
    for t in tables:
        grid = []
        for row in t.rows:
            hs = []
            for c in row.cells:
                got = None
                for para in c.paragraphs:
                    for ext in para._p.iter(qn('wp:extent')):
                        got = int(ext.get('cy')) / EMU
                hs.append(got)
            grid.append(hs)
        if any(any(h is not None for h in hs) for hs in grid):
            n_img = sum(1 for hs in grid for h in hs if h is not None)
            est = sum(
                (max([h for h in hs if h is not None]) if any(h is not None for h in hs) else 0)
                + CAPTION_H + CELL_PAD
                for hs in grid
            )
            photo_pages.append({
                "rows": len(grid), "cols": len(t.columns),
                "n_img": n_img, "est_h": round(est, 2)
            })

    total_img = sum(p["n_img"] for p in photo_pages)
    shape_ok = all(p["rows"] <= MAX_ROWS and p["cols"] == GRID_COLS for p in photo_pages)
    over = [p for p in photo_pages if p["est_h"] > PAGE_BUDGET + 0.01]

    result = {
        "n_tables": len(tables),
        "n_photo_pages": len(photo_pages),
        "n_embedded_imgs": len(imgs),
        "n_imgs_in_tables": total_img,
        "photo_pages": photo_pages,
        "shape_ok": shape_ok,
        "over_budget": len(over),
    }
    if expect_imgs is not None:
        result["expect_match"] = (len(imgs) == expect_imgs and total_img == expect_imgs)
    if verbose:
        print(f"{os.path.basename(path)}")
        print(f"  表格={result['n_tables']} 相片页={result['n_photo_pages']} "
              f"嵌入相={result['n_embedded_imgs']} 表格内相={result['n_imgs_in_tables']}")
        print(f"  每页: {photo_pages}")
        print(f"  形状OK={shape_ok}  超出预算={result['over_budget']}")
        if expect_imgs is not None:
            print(f"  预期匹配={result.get('expect_match')}")
    return result
